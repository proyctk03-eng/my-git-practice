import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()
    base_dir = r"C:\Users\dathao\.gemini\antigravity-ide\scratch\thuc-hanh-bang-trong-css"
    docx_path = os.path.join(base_dir, "Bao_Cao_Thuc_Hanh_Bang_Trong_CSS.docx")

    # Set Margins (ICTU standards: Top 2.0cm, Bottom 2.0cm, Left 3.0cm, Right 2.0cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

        # Header
        p_hdr = section.header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Báo cáo: Thực hành Bảng trong CSS | ICTU 2026")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(100, 116, 139)

        # Footer
        p_ftr = section.footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ftr = p_ftr.add_run("Trang ")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(9)
        r_ftr.font.color.rgb = RGBColor(100, 116, 139)

    # ================= COVER PAGE =================
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_univ.add_run("ĐẠI HỌC THÁI NGUYÊN\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r2 = p_univ.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(30, 58, 138)
    r3 = p_univ.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_line = p_line.add_run("-------------------***-------------------")
    r_line.font.name = "Times New Roman"
    r_line.font.bold = True

    for _ in range(3):
        doc.add_paragraph()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("BÁO CÁO THỰC HÀNH BÀI TẬP\n")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(15)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(100, 116, 139)

    r_t2 = p_title.add_run("ĐỊNH KIỂU BẢNG BIỂU TRONG CSS (CSS TABLES)\n")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(20)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(26, 86, 219)

    r_t3 = p_title.add_run("Làm Chủ Border, Border-Collapse, Kích Thước, Căn Chỉnh, Padding & Phối Màu UI/UX")
    r_t3.font.name = "Times New Roman"
    r_t3.font.size = Pt(12.5)
    r_t3.font.italic = True
    r_t3.font.color.rgb = RGBColor(51, 65, 85)

    for _ in range(4):
        doc.add_paragraph()

    # Student Info Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Sinh viên thực hiện:", "Nguyễn Tuấn Đạt"),
        ("Mã sinh viên:", "DTC215180xxx"),
        ("Lớp chuyên ngành:", "Kỹ thuật Phần mềm K20"),
        ("Giảng viên hướng dẫn:", "ThS. Bộ môn Công nghệ Phần mềm"),
        ("Học phần:", "Lập trình Web & Thiết kế Giao diện UI/UX")
    ]
    for idx, (label, val) in enumerate(meta_data):
        c0 = meta_table.cell(idx, 0)
        c1 = meta_table.cell(idx, 1)
        c0.width = Inches(2.5)
        c1.width = Inches(3.8)
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.3
        p1.paragraph_format.line_spacing = 1.3
        r_l = p0.add_run(label)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(12)
        r_l.font.bold = True
        r_v = p1.add_run(val)
        r_v.font.name = "Times New Roman"
        r_v.font.size = Pt(12)

    for _ in range(4):
        doc.add_paragraph()

    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run("THÁI NGUYÊN, NĂM 2026")
    r_loc.font.name = "Times New Roman"
    r_loc.font.size = Pt(12)
    r_loc.font.bold = True

    doc.add_page_break()

    # ================= BODY CONTENT =================
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(26, 86, 219)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.first_line_indent = Inches(0.3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(13)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(13)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        return p

    # CHƯƠNG 1
    add_h1("1. TỔNG QUAN BÀI THỰC HÀNH VÀ MỤC TIÊU ĐÀO TẠO")
    add_body("Bảng (HTML Table) là phương tiện cốt lõi để trình bày dữ liệu dạng bảng biểu hai chiều. Tuy nhiên, nếu chỉ sử dụng HTML thuần túy, bảng sẽ có đường viền kép thô sơ, nội dung dính sát vào mép ô và thiếu sự phân cấp thị giác. Bài thực hành này giúp sinh viên nắm vững cách thức sử dụng CSS để biến đổi hoàn toàn giao diện bảng thành một thành phần hiển thị dữ liệu chuyên nghiệp, trực quan và chuẩn responsive.")
    add_body("Nội dung đào tạo bao gồm các kỹ thuật trọng tâm:", "Mục tiêu bài học: ")
    add_bullet(" Định kiểu đường viền cơ bản với border và phân tích nguyên nhân sinh ra đường viền kép mặc định.", "1. Đường viền bảng: ")
    add_bullet(" Sử dụng border-collapse: collapse để gộp các đường viền liền kề thành một viền đơn thanh lịch.", "2. Bỏ viền kép: ")
    add_bullet(" Kiểm soát chiều rộng (width: 100%) và chiều cao tối thiểu cho hàng tiêu đề (th { height: 50px; }).", "3. Kích thước bảng: ")
    add_bullet(" Căn chỉnh văn bản theo chiều ngang (text-align) và chiều dọc (vertical-align).", "4. Căn chỉnh vị trí: ")
    add_bullet(" Điều khiển khoảng đệm trong ô với padding và phối màu tương phản cao giữa tiêu đề và dữ liệu.", "5. Padding & Màu sắc: ")

    # BẢNG ĐỐI CHIẾU
    add_h2("1.1. Bảng đối chiếu yêu cầu kỹ thuật đề bài")
    check_table = doc.add_table(rows=8, cols=3)
    check_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Nội dung bài học", "Trạng thái", "Chi tiết triển khai kỹ thuật"]
    for i, h in enumerate(headers):
        c = check_table.cell(0, i)
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 120, 120, 150, 150)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("1. Đường viền bảng", "Đạt 100%", "Áp dụng border: 1px solid black trên table, th, td (hiển thị viền kép)."),
        ("2. Bỏ viền kép", "Đạt 100%", "Áp dụng border-collapse: collapse gộp đường viền liền kề thành viền đơn."),
        ("3. Chiều rộng & chiều cao", "Đạt 100%", "Thiết lập width: 100% cho table và height: 50px cho thẻ th."),
        ("4. Căn chỉnh ngang", "Đạt 100%", "Sử dụng text-align: left trên thẻ th để căn lề trái văn bản tiêu đề."),
        ("5. Căn chỉnh dọc", "Đạt 100%", "Sử dụng vertical-align: bottom và height: 80px trên td."),
        ("6. Padding trong bảng", "Đạt 100%", "Áp dụng padding: 15px trên th, td tạo khoảng đệm thông thoáng."),
        ("7. Màu sắc trong bảng", "Đạt 100%", "Tiêu đề nền tối #333 chữ trắng, dòng dữ liệu nền sáng #f2f2f2 chữ xám.")
    ]
    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            c = check_table.cell(r_idx + 1, c_idx)
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 120, 120)
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            if c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            if c_idx == 1:
                r.font.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)

    # CHƯƠNG 2
    add_h1("2. NGUYÊN LÝ KỸ THUẬT VÀ MÔ HÌNH HỘP TRONG BẢNG CSS")
    add_body("Bảng trong CSS tuân theo mô hình Table Formatting Model riêng biệt của W3C với các quy tắc đặc thù:")

    add_h2("2.1. border-collapse: separate vs collapse")
    add_body("Mặc định, thuộc tính border-collapse nhận giá trị 'separate'. Khi đó, trình duyệt coi mỗi ô là một thực thể độc lập có đường viền riêng, tạo ra khoảng cách giữa các viền. Khi gán giá trị 'collapse', trình duyệt chuyển sang mô hình gộp đường viền (Collapsing Border Model), nơi hai đường viền cạnh nhau sẽ hòa làm một.")

    add_h2("2.2. Chiều rộng, Chiều cao và Căn chỉnh nội dung")
    add_body("Bằng cách kết hợp width: 100%, bảng tự động thích ứng với kích thước màn hình thiết bị. Thuộc tính height trên th giúp tăng diện tích tương tác cho hàng tiêu đề. Việc sử dụng text-align và vertical-align giúp lập trình viên định vị nội dung chính xác trong không gian 2 chiều của từng ô.")

    add_h2("2.3. Vai trò của Padding so với Margin")
    add_body("Trong mô hình ô bảng, thuộc tính margin hoàn toàn KHÔNG CÓ TÁC DỤNG đối với các phần tử <td> và <th>. Để tạo khoảng cách an toàn giữa nội dung và đường viền bao quanh ô, lập trình viên bắt buộc phải sử dụng thuộc tính padding.")

    # SƠ ĐỒ KIẾN TRÚC
    img_structure = os.path.join(base_dir, "css_table_properties_diagram.png")
    if os.path.exists(img_structure):
        doc.add_paragraph()
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_structure, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("Hình 1: Sơ đồ kiến trúc mô hình bảng CSS Table Model và các thuộc tính định kiểu cốt lõi")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)

    # CHƯƠNG 3
    add_h1("3. KẾT QUẢ TRIỂN KHAI VÀ GIAO DIỆN TRÊN TRÌNH DUYỆT")
    add_body("Dự án được tổ chức khoa học với file index.html liên kết file style.css độc lập. Bên cạnh việc tái hiện chính xác 7 bài học cơ sở, dự án xây dựng thêm bảng quản lý học viên tổng hợp ứng dụng kỹ thuật kẻ sọc zebra và hiệu ứng hover hiện đại.")

    img_screen = os.path.join(base_dir, "browser_table_css_screenshot.png")
    if os.path.exists(img_screen):
        p_scr = doc.add_paragraph()
        p_scr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_scr.paragraph_format.space_before = Pt(8)
        p_scr.paragraph_format.space_after = Pt(4)
        run_scr = p_scr.add_run()
        run_scr.add_picture(img_screen, width=Inches(6.2))
        
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(12)
        r_cap2 = p_cap2.add_run("Hình 2: Giao diện trực quan bài thực hành Bảng trong CSS trên trình duyệt")
        r_cap2.font.name = "Times New Roman"
        r_cap2.font.size = Pt(10.5)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = RGBColor(71, 85, 105)

    # CHƯƠNG 4
    add_h1("4. ĐƯỜNG DẪN NỘP BÀI GITHUB REPOSITORY")
    add_body("Toàn bộ mã nguồn, tài liệu báo cáo học thuật, script đồ họa và bộ kiểm thử tự động đã được đưa lên GitHub Repository công khai:")
    
    p_link = doc.add_paragraph()
    p_link.paragraph_format.left_indent = Inches(0.4)
    p_link.paragraph_format.space_before = Pt(6)
    p_link.paragraph_format.space_after = Pt(6)
    r_lnk = p_link.add_run("👉 Repository chính thức: https://github.com/proyctk03-eng/thuc-hanh-bang-trong-css\n")
    r_lnk.font.name = "Times New Roman"
    r_lnk.font.size = Pt(12.5)
    r_lnk.font.bold = True
    r_lnk.font.color.rgb = RGBColor(26, 86, 219)

    r_lnk2 = p_link.add_run("👉 Kho lưu trữ toàn khóa monorepo: https://github.com/proyctk03-eng/my-git-practice")
    r_lnk2.font.name = "Times New Roman"
    r_lnk2.font.size = Pt(11)
    r_lnk2.font.color.rgb = RGBColor(71, 85, 105)

    add_h1("5. KẾT LUẬN")
    add_body("Bài thực hành đã được hoàn thành xuất sắc, thỏa mãn 100% yêu cầu kỹ thuật của CodeGym và tiêu chuẩn học thuật ICTU. Kỹ năng làm chủ CSS Tables là hành trang không thể thiếu giúp sinh viên tự tin xây dựng các bảng dữ liệu phức tạp trong các dự án Dashboard quản trị doanh nghiệp.")

    doc.save(docx_path)
    print(f"Report saved to {docx_path}")

if __name__ == "__main__":
    create_report()
