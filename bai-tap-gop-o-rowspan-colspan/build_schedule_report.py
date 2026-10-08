import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()
    base_dir = r"C:\Users\dathao\.gemini\antigravity-ide\scratch\bai-tap-gop-o-rowspan-colspan"
    docx_path = os.path.join(base_dir, "Bao_Cao_Bai_Tap_Gop_O_Rowspan_Colspan.docx")

    # Set Margins (ICTU standards: Top 2.0cm, Bottom 2.0cm, Left 3.0cm, Right 2.0cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

        # Header
        p_hdr = section.header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Báo cáo: Gộp ô trong bảng HTML với Rowspan và Colspan | ICTU 2026")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(100, 116, 139)

        # Footer
        p_ftr = section.footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ftr = p_ftr.add_run("Trang ")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(9)
        r_ftr.font.color.rgb = RGBColor(100, 116, 139)

    # ================= COVER PAGE =================
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_univ.add_run("ĐẠI HỌC THÁI NGUYÊN\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r2 = p_univ.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(30, 58, 138)
    r3 = p_univ.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_line = p_line.add_run("-------------------***-------------------")
    r_line.font.name = "Times New Roman"
    r_line.font.bold = True

    for _ in range(3):
        doc.add_paragraph()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("BÁO CÁO THỰC HÀNH BÀI TẬP\n")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(15)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(100, 116, 139)

    r_t2 = p_title.add_run("GỘP Ô TRONG BẢNG VỚI ROWSPAN VÀ COLSPAN TRONG HTML5\n")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(20)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(26, 86, 219)

    r_t3 = p_title.add_run("Xây Dựng Giao Diện Bảng Lịch Họp Công Ty Chuẩn Ngữ Nghĩa & Trực Quan Hóa Lưới")
    r_t3.font.name = "Times New Roman"
    r_t3.font.size = Pt(12.5)
    r_t3.font.italic = True
    r_t3.font.color.rgb = RGBColor(51, 65, 85)

    for _ in range(4):
        doc.add_paragraph()

    # Student Info Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Sinh viên thực hiện:", "Nguyễn Tuấn Đạt"),
        ("Mã sinh viên:", "DTC215180xxx"),
        ("Lớp chuyên ngành:", "Kỹ thuật Phần mềm K20"),
        ("Giảng viên hướng dẫn:", "ThS. Bộ môn Công nghệ Phần mềm"),
        ("Học phần:", "Lập trình Web & Thiết kế Giao diện UI/UX")
    ]
    for idx, (label, val) in enumerate(meta_data):
        c0 = meta_table.cell(idx, 0)
        c1 = meta_table.cell(idx, 1)
        c0.width = Inches(2.5)
        c1.width = Inches(3.8)
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.3
        p1.paragraph_format.line_spacing = 1.3
        r_l = p0.add_run(label)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(12)
        r_l.font.bold = True
        r_v = p1.add_run(val)
        r_v.font.name = "Times New Roman"
        r_v.font.size = Pt(12)

    for _ in range(4):
        doc.add_paragraph()

    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run("THÁI NGUYÊN, NĂM 2026")
    r_loc.font.name = "Times New Roman"
    r_loc.font.size = Pt(12)
    r_loc.font.bold = True

    doc.add_page_break()

    # ================= BODY CONTENT =================
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(26, 86, 219)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.first_line_indent = Inches(0.3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(13)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(13)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        return p

    # CHƯƠNG 1
    add_h1("1. TỔNG QUAN BÀI TẬP VÀ MỤC TIÊU KỸ THUẬT")
    add_body("Trong phát triển giao diện web, bảng (HTML Table) là phương thức tiêu chuẩn để cấu trúc hóa dữ liệu dạng bảng biểu hai chiều. Tuy nhiên, các tập dữ liệu thực tế như thời khóa biểu, lịch trực, bảng chấm công hoặc lịch họp doanh nghiệp thường có các sự kiện diễn ra song song hoặc kéo dài qua nhiều mốc thời gian. Khi đó, nếu chỉ dùng các ô đơn lẻ thì bảng sẽ bị lặp dữ liệu, gây rối mắt cho người dùng. Kỹ thuật gộp ô bằng thuộc tính rowspan và colspan giải quyết triệt để bài toán này.")
    add_body("Bài tập yêu cầu người học làm chủ hai kỹ thuật cốt lõi:", "Mục tiêu đào tạo: ")
    add_bullet(" Nắm vững cú pháp và cơ chế hoạt động của thuộc tính gộp dòng theo chiều dọc.", "Thuộc tính rowspan: ")
    add_bullet(" Nắm vững cú pháp và cơ chế hoạt động của thuộc tính gộp cột theo chiều ngang.", "Thuộc tính colspan: ")
    add_bullet(" Khắc phục triệt để lỗi lệch bảng (table broken layout) thông qua nguyên lý bù trừ ô.", "Tư duy lưới toán học: ")

    # BẢNG ĐỐI CHIẾU
    add_h2("1.1. Bảng đối chiếu yêu cầu kỹ thuật đề bài")
    check_table = doc.add_table(rows=6, cols=3)
    check_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Yêu cầu đề bài", "Trạng thái", "Chi tiết giải pháp triển khai"]
    for i, h in enumerate(headers):
        c = check_table.cell(0, i)
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 120, 120, 150, 150)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("Tạo file index.html", "Đạt 100%", "Đã khởi tạo index.html và file đối chiếu index_basic.html."),
        ("Tiêu đề <h1>Lịch họp công ty</h1>", "Đạt 100%", "Tiêu đề h1 đặt tại đầu trang web, chuẩn SEO."),
        ("Bảng gồm cột Ngày, Giờ, Nội dung", "Đạt 100%", "Cấu trúc cột chuẩn xác theo <th> trong <thead>."),
        ("Dùng rowspan khi có nhiều cuộc họp trong ngày", "Đạt 100%", "Áp dụng rowspan=\"2\" (Thứ 2, Thứ 5) và rowspan=\"3\" (Thứ 4)."),
        ("Dùng colspan cho cuộc họp kéo dài nhiều giờ", "Đạt 100%", "Áp dụng colspan=\"2\" cho sự kiện Thứ 3 và Thứ 6 (Cả ngày).")
    ]
    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            c = check_table.cell(r_idx + 1, c_idx)
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 120, 120)
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            if c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            if c_idx == 1:
                r.font.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)

    # CHƯƠNG 2
    add_h1("2. NGUYÊN LÝ LƯỚI & PHÂN TÍCH THUỘC TÍNH ROWSPAN VÀ COLSPAN")
    add_body("Bảng HTML được xây dựng theo mô hình lưới tọa độ (Matrix Grid). Mỗi hàng <tr> đại diện cho một hàng ngang, chứa các ô <td> hoặc <th> tương ứng với các cột. Khi một ô được chỉ định gộp, diện tích hiển thị của nó sẽ mở rộng sang các ô lân cận.")

    add_h2("2.1. Cơ chế hoạt động của Rowspan (Gộp hàng dọc)")
    add_body("Thuộc tính rowspan=\"n\" ra lệnh cho trình duyệt kéo dài ô hiện tại xuống dưới thêm (n - 1) hàng tiếp theo. Ví dụ: khi cột 'Ngày' ở hàng Thứ Hai sử dụng rowspan=\"2\", ô này sẽ chiếm giữ vị trí cột 1 của cả hàng hiện tại và hàng ngay kế tiếp.")
    add_body("Tại hàng <tr> kế tiếp, lập trình viên tuyệt đối KHÔNG được khai báo thẻ <td> cho cột đã bị chiếm chỗ. Nếu khai báo thêm, hàng đó sẽ bị dư ô và đẩy toàn bộ dữ liệu phía sau sang phải, làm vỡ bảng hoàn toàn.", "Quy tắc vàng của Rowspan: ")

    add_h2("2.2. Cơ chế hoạt động của Colspan (Gộp cột ngang)")
    add_body("Thuộc tính colspan=\"m\" mở rộng ô hiện tại sang bên phải thêm (m - 1) cột trên cùng một hàng. Điều này rất thích hợp cho các sự kiện chiếm trọn nhiều mốc giờ hoặc thông báo chung.")
    add_body("Tổng số ô thực tế khai báo trong một hàng (tính cả giá trị colspan) cộng với số ô bị các hàng trên chiếm bởi rowspan PHẢI LUÔN BẰNG tổng số cột của bảng.", "Định luật bảo toàn cột: ")

    # SƠ ĐỒ KIẾN TRÚC
    img_structure = os.path.join(base_dir, "html_rowspan_colspan_diagram.png")
    if os.path.exists(img_structure):
        doc.add_paragraph()
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_structure, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("Hình 1: Sơ đồ cơ chế phân bổ không gian ma trận lưới của Rowspan và Colspan")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)

    # CHƯƠNG 3
    add_h1("3. MÃ NGUỒN TRIỂN KHAI VÀ GIAO DIỆN MINH HỌA")
    add_body("Hệ thống được tổ chức thành 2 tệp nguồn song song: index_basic.html phục vụ việc chấm điểm cấu trúc thuần túy theo giáo trình cốt lõi và index.html phục vụ giao diện chuẩn doanh nghiệp với thiết kế hiện đại, responsive và trực quan hóa tương tác.")

    add_h2("3.1. Mã nguồn cơ bản index_basic.html")
    add_body("Dưới đây là đoạn mã nguồn cô đọng minh chứng đầy đủ các yêu cầu kỹ thuật của đề bài:")

    code_p = doc.add_paragraph()
    code_p.paragraph_format.left_indent = Inches(0.4)
    code_p.paragraph_format.space_before = Pt(4)
    code_p.paragraph_format.space_after = Pt(8)
    code_snippet = (
        '<table border="1" cellpadding="10" cellspacing="0">\n'
        '    <thead>\n'
        '        <tr><th>Ngày</th><th>Giờ</th><th>Nội dung cuộc họp</th></tr>\n'
        '    </thead>\n'
        '    <tbody>\n'
        '        <!-- Thứ Hai: 2 cuộc họp -> rowspan="2" -->\n'
        '        <tr>\n'
        '            <td rowspan="2">Thứ Hai (12/10)</td>\n'
        '            <td>08:30 - 10:00</td>\n'
        '            <td>Họp giao ban đầu tuần toàn công ty</td>\n'
        '        </tr>\n'
        '        <tr>\n'
        '            <td>14:00 - 15:30</td>\n'
        '            <td>Báo cáo tiến độ dự án ERP & Tự động hóa</td>\n'
        '        </tr>\n'
        '        <!-- Thứ Ba: Cuộc họp kéo dài cả ngày -> colspan="2" -->\n'
        '        <tr>\n'
        '            <td>Thứ Ba (13/10)</td>\n'
        '            <td colspan="2">08:00 - 17:00: Hội thảo chiến lược chuyển đổi số (Cả ngày)</td>\n'
        '        </tr>\n'
        '    </tbody>\n'
        '</table>'
    )
    r_code = code_p.add_run(code_snippet)
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(9.5)
    r_code.font.color.rgb = RGBColor(30, 41, 59)

    add_h2("3.2. Ảnh chụp màn hình kết quả chạy trên trình duyệt")
    img_screen = os.path.join(base_dir, "browser_schedule_table_screenshot.png")
    if os.path.exists(img_screen):
        p_scr = doc.add_paragraph()
        p_scr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_scr.paragraph_format.space_before = Pt(8)
        p_scr.paragraph_format.space_after = Pt(4)
        run_scr = p_scr.add_run()
        run_scr.add_picture(img_screen, width=Inches(6.2))
        
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(12)
        r_cap2 = p_cap2.add_run("Hình 2: Giao diện trực quan bảng Lịch họp công ty hiển thị nổi bật ô Rowspan và Colspan")
        r_cap2.font.name = "Times New Roman"
        r_cap2.font.size = Pt(10.5)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = RGBColor(71, 85, 105)

    # CHƯƠNG 4
    add_h1("4. ĐƯỜNG DẪN NỘP BÀI GITHUB REPOSITORY")
    add_body("Toàn bộ mã nguồn bài tập, các script sinh đồ thị, kịch bản kiểm thử tự động cùng báo cáo thuyết minh đã được đưa lên GitHub Repository công khai theo đúng quy định nộp bài của CodeGym:")
    
    p_link = doc.add_paragraph()
    p_link.paragraph_format.left_indent = Inches(0.4)
    p_link.paragraph_format.space_before = Pt(6)
    p_link.paragraph_format.space_after = Pt(6)
    r_lnk = p_link.add_run("👉 Repository chính thức: https://github.com/proyctk03-eng/bai-tap-gop-o-rowspan-colspan\n")
    r_lnk.font.name = "Times New Roman"
    r_lnk.font.size = Pt(12.5)
    r_lnk.font.bold = True
    r_lnk.font.color.rgb = RGBColor(26, 86, 219)

    r_lnk2 = p_link.add_run("👉 Kho lưu trữ toàn khóa: https://github.com/proyctk03-eng/my-git-practice")
    r_lnk2.font.name = "Times New Roman"
    r_lnk2.font.size = Pt(11)
    r_lnk2.font.color.rgb = RGBColor(71, 85, 105)

    add_h1("5. KẾT LUẬN")
    add_body("Bài thực hành đã hoàn thành xuất sắc toàn bộ các mục tiêu đề ra. Việc làm chủ thuộc tính rowspan và colspan không chỉ giúp lập trình viên tạo nên những bảng biểu dữ liệu mạch lạc, khoa học mà còn nâng cao tư duy phân tích lưới cấu trúc dữ liệu - một nền tảng thiết yếu trước khi tiếp cận các công nghệ bố cục hiện đại như CSS Grid và Flexbox.")

    doc.save(docx_path)
    print(f"Report saved to {docx_path}")

if __name__ == "__main__":
    create_report()
