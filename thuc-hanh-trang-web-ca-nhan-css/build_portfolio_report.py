import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()
    base_dir = r"C:\Users\dathao\.gemini\antigravity-ide\scratch\thuc-hanh-trang-web-ca-nhan-css"
    docx_path = os.path.join(base_dir, "Bao_Cao_Thuc_Hanh_Trang_Web_Ca_Nhan_CSS.docx")

    # Set Margins (ICTU standards: Top 2.0cm, Bottom 2.0cm, Left 3.0cm, Right 2.0cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

        # Header
        p_hdr = section.header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Báo cáo: Thực hành Tạo trang web cá nhân với CSS | ICTU 2026")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(100, 116, 139)

        # Footer
        p_ftr = section.footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ftr = p_ftr.add_run("Trang ")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(9)
        r_ftr.font.color.rgb = RGBColor(100, 116, 139)

    # ================= COVER PAGE =================
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_univ.add_run("ĐẠI HỌC THÁI NGUYÊN\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r2 = p_univ.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(30, 58, 138)
    r3 = p_univ.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_line = p_line.add_run("-------------------***-------------------")
    r_line.font.name = "Times New Roman"
    r_line.font.bold = True

    for _ in range(3):
        doc.add_paragraph()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("BÁO CÁO THỰC HÀNH BÀI TẬP\n")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(15)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(100, 116, 139)

    r_t2 = p_title.add_run("TẠO TRANG WEB CÁ NHÂN VỚI CSS (PERSONAL PORTFOLIO)\n")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(20)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(46, 125, 50)

    r_t3 = p_title.add_run("Ứng Dụng Tổng Hợp Bộ Chọn, Đường Viền, Màu Sắc, Ảnh Đại Diện Tròn & Bố Cục Thẩm Mỹ")
    r_t3.font.name = "Times New Roman"
    r_t3.font.size = Pt(12.5)
    r_t3.font.italic = True
    r_t3.font.color.rgb = RGBColor(51, 65, 85)

    for _ in range(4):
        doc.add_paragraph()

    # Student Info Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Sinh viên thực hiện:", "Nguyễn Tuấn Đạt"),
        ("Mã sinh viên:", "DTC215180xxx"),
        ("Lớp chuyên ngành:", "Kỹ thuật Phần mềm K20"),
        ("Giảng viên hướng dẫn:", "ThS. Bộ môn Công nghệ Phần mềm"),
        ("Học phần:", "Lập trình Web & Thiết kế Giao diện UI/UX")
    ]
    for idx, (label, val) in enumerate(meta_data):
        c0 = meta_table.cell(idx, 0)
        c1 = meta_table.cell(idx, 1)
        c0.width = Inches(2.5)
        c1.width = Inches(3.8)
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.3
        p1.paragraph_format.line_spacing = 1.3
        r_l = p0.add_run(label)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(12)
        r_l.font.bold = True
        r_v = p1.add_run(val)
        r_v.font.name = "Times New Roman"
        r_v.font.size = Pt(12)

    for _ in range(4):
        doc.add_paragraph()

    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run("THÁI NGUYÊN, NĂM 2026")
    r_loc.font.name = "Times New Roman"
    r_loc.font.size = Pt(12)
    r_loc.font.bold = True

    doc.add_page_break()

    # ================= BODY CONTENT =================
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(46, 125, 50)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.first_line_indent = Inches(0.3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(13)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(13)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        return p

    # CHƯƠNG 1
    add_h1("1. TỔNG QUAN BÀI THỰC HÀNH VÀ MỤC TIÊU KỸ THUẬT")
    add_body("Trang web cá nhân (Personal Profile / Portfolio) là dự án khởi đầu kinh điển trong lộ trình học phát triển web Front-end. Mục tiêu cốt lõi của bài thực hành là tích hợp đồng thời nhiều mảng kiến thức CSS đã học vào một trang web hoàn chỉnh, bao gồm: cấu trúc ngữ nghĩa HTML5, nhúng CSS ngoại vi, vận dụng linh hoạt các bộ chọn phần tử và bộ chọn lớp, căn chỉnh bố cục trung tâm, xử lý hình ảnh và định dạng bảng biểu thông tin liên hệ.")
    add_body("Mục tiêu học tập cụ thể:", "Năng lực đạt được: ")
    add_bullet(" Liên kết tệp CSS độc lập style.css vào tài liệu HTML qua thẻ <link>.", "1. Nhúng CSS: ")
    add_bullet(" Sử dụng thành thạo Type Selector (body, header, ul, table...) và Class Selector (.avatar, .profile, .hobbies...).", "2. Bộ chọn CSS: ")
    add_bullet(" Định dạng đường viền (border) và bo tròn hoàn hảo ảnh đại diện (border-radius: 50%).", "3. Xử lý đường viền: ")
    add_bullet(" Phối hợp màu nền background-color và ảnh nền lặp nhẹ nhàng background-image.", "4. Quản lý nền: ")
    add_bullet(" Bố cục bảng liên hệ gọn gàng với width: 50%, margin: auto và border-collapse: collapse.", "5. Bảng biểu thông tin: ")

    # BẢNG ĐỐI CHIẾU
    add_h2("1.1. Bảng đối chiếu yêu cầu kỹ thuật đề bài")
    check_table = doc.add_table(rows=7, cols=3)
    check_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Yêu cầu thành phần", "Trạng thái", "Chi tiết giải pháp triển khai"]
    for i, h in enumerate(headers):
        c = check_table.cell(0, i)
        set_cell_background(c, "2E7D32")
        set_cell_margins(c, 120, 120, 150, 150)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("Tiêu đề chính", "Đạt 100%", "Thẻ <header> chứa <h1> với màu nền xanh lá #4CAF50 chữ trắng."),
        ("Ảnh đại diện", "Đạt 100%", "File avatar.jpg, class .avatar với border-radius: 50% và border 3px."),
        ("Giới thiệu bản thân", "Đạt 100%", "<section class=\"profile\"> chứa thông tin họ tên, nghề nghiệp."),
        ("Danh sách sở thích", "Đạt 100%", "<section class=\"hobbies\"> thẻ <ul> với list-style-type: square."),
        ("Mục tiêu học tập", "Đạt 100%", "<section class=\"goals\"> chứa mục tiêu phát triển nghề nghiệp."),
        ("Bảng thông tin liên hệ", "Đạt 100%", "<footer> chứa <table> căn giữa (margin: auto) và border-collapse.")
    ]
    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            c = check_table.cell(r_idx + 1, c_idx)
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 120, 120)
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            if c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            if c_idx == 1:
                r.font.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)

    # CHƯƠNG 2
    add_h1("2. NGUYÊN LÝ KỸ THUẬT VÀ PHÂN TÍCH ĐỊNH KIỂU CSS")
    add_body("Trang web cá nhân khai thác các kỹ thuật tạo hình thị giác hiện đại của CSS3:")

    add_h2("2.1. Kỹ thuật tạo ảnh tròn với border-radius: 50%")
    add_body("Để biến một bức ảnh vuông thành hình tròn hoàn hảo, kích thước chiều rộng (width) và chiều cao (height) phải bằng nhau tuyệt đối (150px x 150px). Khi áp dụng border-radius: 50%, trình duyệt sẽ vẽ bán kính bo tròn bằng một nửa chiều rộng của ảnh. Kết hợp với thuộc tính border: 3px solid #4CAF50 tạo nên đường viền nổi bật đóng khung ảnh đại diện.")

    add_h2("2.2. Phối hợp background-color và background-image")
    add_body("CSS cho phép khai báo đồng thời màu nền và ảnh nền. Màu nền background-color: #f5f5f5 đóng vai trò làm lớp màu cơ sở an toàn (fallback), trong khi background-image: url('bg_pattern.png') bổ sung họa tiết hình học tinh tế, giúp trang web không bị đơn điệu.")

    add_h2("2.3. Bố cục căn giữa danh sách và bảng biểu")
    add_body("Đối với danh sách sở thích <ul>, việc áp dụng display: inline-block kết hợp text-align: left giúp toàn bộ khối danh sách nằm gọn giữa trang web mà các mục <li> bên trong vẫn thẳng hàng bên trái. Đối với bảng thông tin liên hệ, margin: auto kết hợp width: 50% là giải pháp căn giữa kinh điển và chuẩn hóa nhất.")

    # SƠ ĐỒ KIẾN TRÚC
    img_structure = os.path.join(base_dir, "css_portfolio_structure_diagram.png")
    if os.path.exists(img_structure):
        doc.add_paragraph()
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_structure, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("Hình 1: Sơ đồ kiến trúc cấu trúc phân cấp DOM và các bộ chọn định kiểu CSS của trang web cá nhân")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)

    # CHƯƠNG 3
    add_h1("3. KẾT QUẢ TRIỂN KHAI VÀ GIAO DIỆN TRÊN TRÌNH DUYỆT")
    add_body("Mã nguồn dự án tuân thủ nghiêm ngặt nguyên tắc phân tách giữa HTML và CSS. Giao diện trang web hài hòa với tông màu xanh lá chủ đạo (#4CAF50 và #2e7d32), ảnh avatar tròn sắc nét, danh sách sở thích bullet vuông và bảng liên hệ cân đối.")

    img_screen = os.path.join(base_dir, "browser_portfolio_screenshot.png")
    if os.path.exists(img_screen):
        p_scr = doc.add_paragraph()
        p_scr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_scr.paragraph_format.space_before = Pt(8)
        p_scr.paragraph_format.space_after = Pt(4)
        run_scr = p_scr.add_run()
        run_scr.add_picture(img_screen, width=Inches(6.2))
        
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(12)
        r_cap2 = p_cap2.add_run("Hình 2: Ảnh chụp màn hình kết quả chạy trang web cá nhân trên trình duyệt")
        r_cap2.font.name = "Times New Roman"
        r_cap2.font.size = Pt(10.5)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = RGBColor(71, 85, 105)

    # CHƯƠNG 4
    add_h1("4. ĐƯỜNG DẪN NỘP BÀI GITHUB REPOSITORY")
    add_body("Mã nguồn hoàn thiện, các tệp đồ họa, ảnh avatar, script kiểm thử tự động và báo cáo thuyết minh đã được đưa lên GitHub Repository công khai theo đúng quy định nộp bài của CodeGym:")
    
    p_link = doc.add_paragraph()
    p_link.paragraph_format.left_indent = Inches(0.4)
    p_link.paragraph_format.space_before = Pt(6)
    p_link.paragraph_format.space_after = Pt(6)
    r_lnk = p_link.add_run("👉 Repository chính thức: https://github.com/proyctk03-eng/thuc-hanh-trang-web-ca-nhan-css\n")
    r_lnk.font.name = "Times New Roman"
    r_lnk.font.size = Pt(12.5)
    r_lnk.font.bold = True
    r_lnk.font.color.rgb = RGBColor(46, 125, 50)

    r_lnk2 = p_link.add_run("👉 Kho lưu trữ toàn khóa monorepo: https://github.com/proyctk03-eng/my-git-practice")
    r_lnk2.font.name = "Times New Roman"
    r_lnk2.font.size = Pt(11)
    r_lnk2.font.color.rgb = RGBColor(71, 85, 105)

    add_h1("5. KẾT LUẬN")
    add_body("Bài thực hành đã được hoàn thành trọn vẹn 100%. Thông qua việc xây dựng trang web cá nhân, sinh viên đã củng cố vững chắc năng lực phối hợp các thuộc tính CSS cơ bản (font chữ, màu sắc, ảnh nền, đường viền, danh sách và bảng), tạo tiền đề vững chắc cho các bài tập bố cục nâng cao tiếp theo.")

    doc.save(docx_path)
    print(f"Report saved to {docx_path}")

if __name__ == "__main__":
    create_report()
