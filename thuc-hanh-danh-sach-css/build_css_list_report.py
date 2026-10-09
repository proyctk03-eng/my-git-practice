import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()
    base_dir = r"C:\Users\dathao\.gemini\antigravity-ide\scratch\thuc-hanh-danh-sach-css"
    docx_path = os.path.join(base_dir, "Bao_Cao_Thuc_Hanh_Danh_Sach_Trong_CSS.docx")

    # Set Margins (ICTU standards: Top 2.0cm, Bottom 2.0cm, Left 3.0cm, Right 2.0cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

        # Header
        p_hdr = section.header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Báo cáo: Thực hành Danh sách trong CSS | ICTU 2026")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(100, 116, 139)

        # Footer
        p_ftr = section.footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ftr = p_ftr.add_run("Trang ")
        r_ftr.font.name = "Times New Roman"
        r_ftr.font.size = Pt(9)
        r_ftr.font.color.rgb = RGBColor(100, 116, 139)

    # ================= COVER PAGE =================
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_univ.add_run("ĐẠI HỌC THÁI NGUYÊN\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r2 = p_univ.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(30, 58, 138)
    r3 = p_univ.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_line = p_line.add_run("-------------------***-------------------")
    r_line.font.name = "Times New Roman"
    r_line.font.bold = True

    for _ in range(3):
        doc.add_paragraph()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("BÁO CÁO THỰC HÀNH BÀI TẬP\n")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(15)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(100, 116, 139)

    r_t2 = p_title.add_run("ĐỊNH KIỂU DANH SÁCH TRONG CSS (CSS LISTS)\n")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(20)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(26, 86, 219)

    r_t3 = p_title.add_run("Làm Chủ Các Thuộc Tính List-Style, Marker Định Kiểu & Xây Dựng Lộ Trình Web Hiện Đại")
    r_t3.font.name = "Times New Roman"
    r_t3.font.size = Pt(12.5)
    r_t3.font.italic = True
    r_t3.font.color.rgb = RGBColor(51, 65, 85)

    for _ in range(4):
        doc.add_paragraph()

    # Student Info Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Sinh viên thực hiện:", "Nguyễn Tuấn Đạt"),
        ("Mã sinh viên:", "DTC215180xxx"),
        ("Lớp chuyên ngành:", "Kỹ thuật Phần mềm K20"),
        ("Giảng viên hướng dẫn:", "ThS. Bộ môn Công nghệ Phần mềm"),
        ("Học phần:", "Lập trình Web & Thiết kế Giao diện UI/UX")
    ]
    for idx, (label, val) in enumerate(meta_data):
        c0 = meta_table.cell(idx, 0)
        c1 = meta_table.cell(idx, 1)
        c0.width = Inches(2.5)
        c1.width = Inches(3.8)
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.3
        p1.paragraph_format.line_spacing = 1.3
        r_l = p0.add_run(label)
        r_l.font.name = "Times New Roman"
        r_l.font.size = Pt(12)
        r_l.font.bold = True
        r_v = p1.add_run(val)
        r_v.font.name = "Times New Roman"
        r_v.font.size = Pt(12)

    for _ in range(4):
        doc.add_paragraph()

    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run("THÁI NGUYÊN, NĂM 2026")
    r_loc.font.name = "Times New Roman"
    r_loc.font.size = Pt(12)
    r_loc.font.bold = True

    doc.add_page_break()

    # ================= BODY CONTENT =================
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(26, 86, 219)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.first_line_indent = Inches(0.3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(13)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Times New Roman"
            r_pre.font.size = Pt(13)
            r_pre.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        return p

    # CHƯƠNG 1
    add_h1("1. TỔNG QUAN BÀI THỰC HÀNH VÀ MỤC TIÊU ĐÀO TẠO")
    add_body("Danh sách (Lists) là một trong những thành phần cơ bản và quan trọng nhất trong việc tổ chức nội dung có thứ tự hoặc không có thứ tự trên các trang web. Mặc định, trình duyệt cung cấp các kiểu đánh dấu đơn giản như dấu chấm tròn (disc) hoặc số đếm thập phân (1, 2, 3). Để xây dựng giao diện hiện đại, chuyên nghiệp và đồng bộ theo nhận diện thương hiệu, lập trình viên cần nắm vững nhóm thuộc tính CSS List Properties và mối liên hệ với mô hình hộp (Box Model).")
    add_body("Bài thực hành hướng đến việc trang bị các kỹ năng cốt lõi:", "Mục tiêu bài học: ")
    add_bullet(" Sử dụng list-style-type để thay đổi kiểu đánh dấu sang dạng hình vuông (square) hoặc số La Mã viết hoa (upper-roman).", "1. Kiểu đánh dấu: ")
    add_bullet(" Sử dụng list-style-image để tải biểu tượng bullet tùy biến dạng hình ảnh hoặc icon.", "2. Hình ảnh đánh dấu: ")
    add_bullet(" Phân biệt rõ sự khác nhau giữa list-style-position: outside (dấu nằm ngoài) và inside (dấu nằm trong).", "3. Vị trí đánh dấu: ")
    add_bullet(" Tối ưu hóa cú pháp CSS bằng thuộc tính rút gọn list-style.", "4. Thuộc tính rút gọn: ")
    add_bullet(" Kết hợp hài hòa màu sắc (background-color, color) và căn chỉnh khoảng cách (padding, margin).", "5. Thiết kế nâng cao: ")

    # BẢNG ĐỐI CHIẾU
    add_h2("1.1. Bảng đối chiếu yêu cầu kỹ thuật đề bài")
    check_table = doc.add_table(rows=9, cols=3)
    check_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Nội dung yêu cầu", "Trạng thái", "Chi tiết triển khai kỹ thuật"]
    for i, h in enumerate(headers):
        c = check_table.cell(0, i)
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 120, 120, 150, 150)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("Mục 1: Chuẩn bị file", "Đạt 100%", "Tạo index.html và style.css, liên kết chuẩn qua thẻ <link>."),
        ("Mục 2: Kiểu đánh dấu square", "Đạt 100%", "Áp dụng list-style-type: square cho 4 môn học: HTML, CSS, JS, Git."),
        ("Mục 3: Đánh số upper-roman", "Đạt 100%", "Áp dụng list-style-type: upper-roman cho danh sách bước học <ol>."),
        ("Mục 4: list-style-image", "Đạt 100%", "Tạo icon bullet.png cục bộ và liên kết qua hàm url('bullet.png')."),
        ("Mục 5: outside vs inside", "Đạt 100%", "Tạo 2 danh sách so sánh trực quan với viền đứt nét để quan sát vị trí."),
        ("Mục 6: list-style rút gọn", "Đạt 100%", "Kết hợp: list-style: square inside url('bullet.png')."),
        ("Mục 7: Màu sắc & khoảng cách", "Đạt 100%", "Áp dụng background-color, color, padding, margin cho danh sách ngôn ngữ."),
        ("Mục 8: Bài tập tổng hợp", "Đạt 100%", "Xây dựng Lộ trình học lập trình Web (5 chặng) chuẩn UI/UX và responsive.")
    ]
    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            c = check_table.cell(r_idx + 1, c_idx)
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 120, 120)
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            if c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            if c_idx == 1:
                r.font.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)

    # CHƯƠNG 2
    add_h1("2. NGUYÊN LÝ KỸ THUẬT VÀ PHÂN TÍCH THUỘC TÍNH CSS LIST")
    add_body("Nhóm thuộc tính định kiểu danh sách trong CSS bao gồm 3 thuộc tính thành phần và 1 thuộc tính rút gọn:")

    add_h2("2.1. list-style-type và các hệ thống đánh dấu")
    add_body("Thuộc tính list-style-type quy định biểu tượng hình học hoặc ký tự đếm của thẻ <li>. Đối với danh sách không thứ tự (ul), các giá trị phổ biến là disc, circle, square và none. Đối với danh sách có thứ tự (ol), CSS hỗ trợ hệ thống số học phong phú như decimal, upper-roman, lower-roman, upper-alpha, lower-alpha.")

    add_h2("2.2. list-style-position: outside so với inside")
    add_body("Đây là thuộc tính dễ gây hiểu nhầm nhất cho người mới bắt đầu lập trình:")
    add_bullet(" Giá trị mặc định của trình duyệt. Dấu đánh dấu nằm hoàn toàn bên ngoài lề nội dung của phần tử <li>. Khi đoạn văn bản trong mục danh sách dài hơn 1 dòng, các dòng tiếp theo sẽ thẳng hàng với dòng đầu tiên.", "list-style-position: outside: ")
    add_bullet(" Dấu đánh dấu được kéo vào nằm bên trong hộp nội dung của <li> và được xử lý như một ký tự văn bản thông thường. Khi văn bản xuống dòng, dòng thứ hai sẽ thụt lề về phía bên trái ngay dưới dấu đánh dấu.", "list-style-position: inside: ")

    add_h2("2.3. list-style-image và thuộc tính rút gọn list-style")
    add_body("Thuộc tính list-style-image: url(...) cho phép nhà phát triển sử dụng ảnh biểu tượng. Thuộc tính rút gọn list-style gom 3 thuộc tính theo thứ tự: list-style: [type] [position] [image]. Trình duyệt sẽ ưu tiên hiển thị hình ảnh nếu có, và tự động fallback về list-style-type nếu ảnh bị lỗi đường dẫn.")

    # SƠ ĐỒ KIẾN TRÚC
    img_structure = os.path.join(base_dir, "css_list_properties_diagram.png")
    if os.path.exists(img_structure):
        doc.add_paragraph()
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_structure, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("Hình 1: Sơ đồ kiến trúc và cơ chế phân cấp các thuộc tính CSS List theo chuẩn W3C")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)

    # CHƯƠNG 3
    add_h1("3. KẾT QUẢ TRIỂN KHAI VÀ GIAO DIỆN TRÊN TRÌNH DUYỆT")
    add_body("Mã nguồn được tổ chức tách bạch hoàn toàn giữa cấu trúc (index.html) và định kiểu (style.css). Giao diện thể hiện rõ nét từng mục yêu cầu với các hộp thẻ trực quan và thẻ tổng hợp 'Lộ trình học lập trình Web' hiện đại.")

    img_screen = os.path.join(base_dir, "browser_css_list_screenshot.png")
    if os.path.exists(img_screen):
        p_scr = doc.add_paragraph()
        p_scr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_scr.paragraph_format.space_before = Pt(8)
        p_scr.paragraph_format.space_after = Pt(4)
        run_scr = p_scr.add_run()
        run_scr.add_picture(img_screen, width=Inches(6.2))
        
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(12)
        r_cap2 = p_cap2.add_run("Hình 2: Ảnh chụp màn hình kết quả chạy bài thực hành Danh sách trong CSS trên trình duyệt")
        r_cap2.font.name = "Times New Roman"
        r_cap2.font.size = Pt(10.5)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = RGBColor(71, 85, 105)

    # CHƯƠNG 4
    add_h1("4. ĐƯỜNG DẪN NỘP BÀI GITHUB REPOSITORY")
    add_body("Mã nguồn dự án, tệp định kiểu, ảnh biểu tượng, script kiểm thử tự động và báo cáo thuyết minh đã được đưa lên GitHub Repository công khai theo đúng quy định nộp bài CodeGym:")
    
    p_link = doc.add_paragraph()
    p_link.paragraph_format.left_indent = Inches(0.4)
    p_link.paragraph_format.space_before = Pt(6)
    p_link.paragraph_format.space_after = Pt(6)
    r_lnk = p_link.add_run("👉 Repository chính thức: https://github.com/proyctk03-eng/thuc-hanh-danh-sach-css\n")
    r_lnk.font.name = "Times New Roman"
    r_lnk.font.size = Pt(12.5)
    r_lnk.font.bold = True
    r_lnk.font.color.rgb = RGBColor(26, 86, 219)

    r_lnk2 = p_link.add_run("👉 Kho lưu trữ toàn khóa monorepo: https://github.com/proyctk03-eng/my-git-practice")
    r_lnk2.font.name = "Times New Roman"
    r_lnk2.font.size = Pt(11)
    r_lnk2.font.color.rgb = RGBColor(71, 85, 105)

    add_h1("5. KẾT LUẬN")
    add_body("Bài thực hành đã hoàn thành trọn vẹn 100% mục tiêu đề ra. Việc nắm vững cách phối hợp giữa list-style, màu sắc và mô hình hộp giúp tạo nên những thành phần điều hướng (Navigation Bar), danh sách tính năng (Feature List) hoặc các bước quy trình (Stepper UI) thanh lịch, nâng cao trải nghiệm người dùng trên các sản phẩm phần mềm thực tế.")

    doc.save(docx_path)
    print(f"Report saved to {docx_path}")

if __name__ == "__main__":
    create_report()
